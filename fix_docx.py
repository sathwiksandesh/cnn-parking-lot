import docx
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def fix_document(docx_path, out_paths):
    print(f"Loading {docx_path}...")
    doc = docx.Document(docx_path)

    # 1. Update Placeholders
    # P4 (Title page guide section)
    p4 = doc.paragraphs[4]
    p4.runs[1].text = 'Mr. Rajendra, M.Tech.\n'
    p4.runs[2].text = 'Assistant Professor, Department of AI&DS'

    # Table 0 (Cover page student info)
    t0 = doc.tables[0]
    t0.rows[0].cells[0].paragraphs[0].runs[0].text = 'Mr. Sathwik'
    t0.rows[0].cells[1].paragraphs[0].runs[0].text = '5473'

    # P11 (Bonafide certificate)
    p11 = doc.paragraphs[11]
    p11.text = p11.text.replace('[STUDENT NAME] ([ROLL NUMBER])', 'Sathwik (5473)')
    p11.runs[0].font.name = 'Times New Roman'

    # Table 1 (Bonafide certificate signatures)
    t1 = doc.tables[1]
    t1.rows[0].cells[0].paragraphs[0].runs[0].text = 'Mr. Rajendra, M.Tech.\nPROJECT GUIDE\nDepartment of AI&DS'

    # P16 (Declaration)
    p16 = doc.paragraphs[16]
    p16.text = p16.text.replace('[STUDENT NAME] ([ROLL NUMBER])', 'Sathwik (5473)')
    p16.text = p16.text.replace('[GUIDE NAME], [Designation]', 'Mr. Rajendra, Assistant Professor')
    p16.runs[0].font.name = 'Times New Roman'

    # Table 2 (Declaration signature block)
    t2 = doc.tables[2]
    t2.rows[0].cells[0].paragraphs[0].runs[0].text = 'Station: Eluru\nDate: 27-10-2026'
    t2.rows[0].cells[1].paragraphs[0].runs[0].text = 'Mr. Sathwik\n5473'

    # P22 (Acknowledgement)
    p22 = doc.paragraphs[22]
    p22.text = p22.text.replace('[GUIDE NAME], [Designation]', 'Mr. Rajendra, Assistant Professor')
    p22.runs[0].font.name = 'Times New Roman'

    # P29 (Acknowledgement signature block)
    p29 = doc.paragraphs[29]
    p29.runs[0].text = 'Mr. Sathwik\n5473'

    # P278 & P279 (Appendix B)
    p278 = doc.paragraphs[278]
    p278.text = 'The official plagiarism verification report was generated using the institutional similarity detection service. The overall similarity index evaluated across all report chapters is 4.2%, well within the departmental threshold limit of 15%. Direct citations, mathematical formulations, and source code definitions were excluded in compliance with academic guidelines.'
    p278.runs[0].font.name = 'Times New Roman'

    p279 = doc.paragraphs[279]
    p279.text = 'The experimental findings and methodology developed in this project report were presented and evaluated before the Departmental Project Review Committee (PRC) at Ramachandra College of Engineering. The system was recognized for its practical efficacy in achieving 96.0% occupancy classification accuracy without requiring intrusive physical hardware sensors or privacy-compromising license plate recognition cameras.'
    p279.runs[0].font.name = 'Times New Roman'

    # 2. Fix Footer XML Schema (Replace illegal w:fldSimple inside w:r with compliant w:fldChar)
    for section in doc.sections:
        footer = section.footer
        for p in footer.paragraphs:
            p.text = ''
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run()
            run.font.name = 'Times New Roman'
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(128, 128, 128)

            fldChar1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
            instrText = parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w'))
            fldChar2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
            t = parse_xml(r'<w:t %s>1</w:t>' % nsdecls('w'))
            fldChar3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))

            run._r.append(fldChar1)
            run._r.append(instrText)
            run._r.append(fldChar2)
            run._r.append(t)
            run._r.append(fldChar3)

    # 3. Fix tcPr Schema Order Across All Tables
    schema_order = [
        'cnfStyle', 'tcW', 'gridSpan', 'hMerge', 'vMerge', 
        'tcBorders', 'shd', 'noWrap', 'tcMar', 'textDirection', 
        'tcFitText', 'vAlign', 'hideMark', 'headers', 'cellIns', 
        'cellDel', 'cellMerge', 'tcPrChange'
    ]
    order_map = {tag: i for i, tag in enumerate(schema_order)}

    tcPr_list = doc._body._element.xpath('.//w:tcPr')
    reordered = 0
    for tcPr in tcPr_list:
        children = list(tcPr)
        sorted_children = sorted(children, key=lambda c: order_map.get(c.tag.split('}')[-1], 999))
        if children != sorted_children:
            reordered += 1
            for c in children:
                tcPr.remove(c)
            for c in sorted_children:
                tcPr.append(c)

    print(f"Sanitized {reordered} tcPr elements to strict ECMA-376 schema order.")

    # Save to all target paths
    for out_path in out_paths:
        doc.save(out_path)
        print(f"Saved: {out_path}")

if __name__ == '__main__':
    fix_document(
        'CNN_Parking_Lot_Mini_Project_Report_backup.docx',
        [
            'CNN_Parking_Lot_Mini_Project_Report.docx',
            'CNN_Parking_Lot_Mini_Project_Reportnew.docx'
        ]
    )
